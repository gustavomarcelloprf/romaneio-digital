# --- Bibliotecas Padrão do Python ---
import os
import re
import io
import csv
import hmac
import unicodedata
import functools
import secrets
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
import sqlite3
import sys
sys.path.append(str(Path(__file__).resolve().parent))

# --- Bibliotecas de Terceiros (Flask, etc.) ---
from flask import (
    Flask, request, jsonify, send_from_directory, send_file,
    session, abort, redirect, url_for, render_template
)
from werkzeug.security import generate_password_hash, check_password_hash
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address
from openpyxl import load_workbook
import openpyxl

# --- Módulos do Projeto (Romaneio Digital) ---
from database import (
    init_db, get_conn, salvar_pix_by_codigo,
    get_pedido, get_itens_pedido, get_loja_by_codigo,
    get_usuario, get_usuario_by_id
)
from services.render_png import render_pedido_png
from services.render_pdf import render_pedido_pdf

# -----------------------------------------------------------------------------
# Configuração do App e Paths
# -----------------------------------------------------------------------------
PROJECT_ROOT = Path(__file__).resolve().parent.parent
FRONT_DIR = PROJECT_ROOT / "frontend"
HTML_DIR = FRONT_DIR / "html"

app = Flask(__name__, template_folder='templates')
app.config.update(
    JSON_AS_ASCII=False,
    SECRET_KEY=os.environ["SECRET_KEY"],
    TEMPLATES_AUTO_RELOAD=True,
    SESSION_COOKIE_SECURE=os.environ.get("SESSION_COOKIE_SECURE", "true").lower() != "false",
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
)
limiter = Limiter(get_remote_address, app=app, storage_uri="memory://")

if not os.environ.get("ADMIN_TOKEN"):
    raise RuntimeError("ADMIN_TOKEN não configurado.")

init_db()

# -----------------------------------------------------------------------------
# Papéis e Permissões
# -----------------------------------------------------------------------------
# Mapa central: capacidade -> papéis que a possuem. Os decorators e o /me
# consultam SÓ este mapa, então encaixar um papel novo (ex.: 'gerente' ganhar
# mais poderes) é editar uma linha aqui — nenhuma rota precisa mudar.
#
# 'gerente' já está desenhado e valendo no backend; ainda não há UI para criar
# um (a criação de usuário continua gerando 'operador').
PAPEIS = ("admin", "gerente", "operador")

PERMISSOES = {
    # Vender / criar pedido: todo mundo vende.
    "vender":              ("operador", "gerente", "admin"),
    "pedidos_ver":         ("operador", "gerente", "admin"),
    "clientes_gerir":      ("operador", "gerente", "admin"),
    # Fiado: vender fiado é dentro do /pedidos (liberado a quem vende); ver
    # saldos e registrar pagamento exige esta capacidade.
    "fiado_gerir":         ("gerente", "admin"),
    # Estoque: qualquer um consulta; só gerente/admin mexem
    # (inclui a entrada de mercadoria / recebimento).
    "estoque_ver":         ("operador", "gerente", "admin"),
    "estoque_mutar":       ("gerente", "admin"),
    # Dashboard da loja: gerente vê o desempenho; só o admin (dono) vê
    # quanto tem de comissão a PAGAR.
    "relatorio_loja":      ("gerente", "admin"),
    "relatorio_comissoes": ("admin",),
    # Despesas (importar/listar/remover): dinheiro que sai do bolso do dono.
    "despesas_gerir":      ("admin",),
    # Gestão de gente e configuração da loja: dono apenas.
    "usuarios_gerir":      ("admin",),
    "config_editar":       ("admin",),
}


def papel_atual():
    return session.get("papel")


def pode(permissao, papel=None):
    """True se o papel informado (ou o da sessão) tem a permissão."""
    if papel is None:
        papel = papel_atual()
    return papel in PERMISSOES.get(permissao, ())


def permissoes_do_papel(papel):
    """Dicionário capacidade -> bool, consumido pelo front para montar o menu."""
    return {nome: (papel in papeis) for nome, papeis in PERMISSOES.items()}


# -----------------------------------------------------------------------------
# Funções Helper
# -----------------------------------------------------------------------------
def _ptbr_to_float(val):
    if val is None: return None
    if isinstance(val, (int, float)): return float(val)
    s = re.sub(r"[^\d,.\-]", "", str(val).strip()).replace(".", "").replace(",", ".")
    try: return float(s)
    except (ValueError, TypeError): return None

class EstoqueInsuficiente(Exception):
    """Venda que o estoque não cobre: aborta a transação com mensagem clara (409)."""

def _baixar_estoque(conn, loja, tecido_nome, itens):
    """Baixa POR PESO do estoque, dentro da transação do chamador.

    `itens` são pares (cor, peso). A venda sai do total (kg) da cor; várias
    linhas da mesma cor somam antes de validar. qtd_pecas é só informativo e
    não é decrementado. Levanta EstoqueInsuficiente (-> 409) se não cobrir.
    """
    tecido_row = conn.execute("SELECT id FROM estoque_tecidos WHERE nome_tecido = ? AND loja = ?", (tecido_nome, loja)).fetchone()
    if not tecido_row:
        raise EstoqueInsuficiente(f"Tipo de tecido '{tecido_nome}' não encontrado no estoque.")
    tecido_id = tecido_row['id']

    peso_por_cor = {}
    for cor, peso in itens:
        cor_nome = (cor or "").strip()
        peso_pedido = peso or 0
        if not cor_nome or peso_pedido <= 0: continue
        peso_por_cor[cor_nome] = peso_por_cor.get(cor_nome, 0) + peso_pedido

    for cor_nome, peso_pedido in peso_por_cor.items():
        cor_row = conn.execute("SELECT id, peso_kg FROM estoque_cores WHERE tecido_id = ? AND nome_cor = ?", (tecido_id, cor_nome)).fetchone()
        if not cor_row:
            raise EstoqueInsuficiente(f"Cor '{cor_nome}' não encontrada no estoque de '{tecido_nome}'.")
        if cor_row['peso_kg'] + 1e-9 < peso_pedido:
            raise EstoqueInsuficiente(
                f"Peso insuficiente para {tecido_nome} - cor {cor_nome}: "
                f"pedido {round(peso_pedido, 3)} kg, disponível {round(cor_row['peso_kg'], 3)} kg."
            )
        # ROUND evita resíduo de ponto flutuante (ex.: 0,30000000004 kg).
        conn.execute("UPDATE estoque_cores SET peso_kg = ROUND(peso_kg - ?, 3) WHERE id = ?", (peso_pedido, cor_row['id']))

def _taxa_comissao(usuario):
    """Taxa vigente do vendedor. O admin é o dono e não tira comissão de si mesmo."""
    if not usuario or usuario.get("papel") == "admin":
        return 0.0
    return float(usuario.get("taxa_comissao") or 0)

def current_loja():
    return (session.get("loja_codigo") or "").strip()

def require_login(func):
    @functools.wraps(func)
    def wrapper(*args, **kwargs):
        loja_codigo = current_loja()
        usuario_id = session.get("usuario_id")
        if not loja_codigo or not usuario_id:
            return jsonify(error="Não autenticado. Faça login primeiro."), 401
        loja = get_loja_by_codigo(loja_codigo)
        if not loja or loja.get("status") != "aprovado":
            session.clear()
            return jsonify(error="Sessão inválida. Faça login novamente."), 401
        usuario = get_usuario_by_id(usuario_id)
        if not usuario or usuario.get("loja") != loja_codigo or not usuario.get("ativo"):
            session.clear()
            return jsonify(error="Sessão inválida. Faça login novamente."), 401
        return func(*args, **kwargs)
    return wrapper

def require_perm(permissao):
    """Exige uma CAPACIDADE, não um papel: a lista de papéis vive em PERMISSOES."""
    def decorator(func):
        @functools.wraps(func)
        def wrapper(*args, **kwargs):
            if not pode(permissao):
                return jsonify(error="Acesso restrito: permissão insuficiente."), 403
            return func(*args, **kwargs)
        return wrapper
    return decorator

def require_admin(func):
    @functools.wraps(func)
    def wrapper(*args, **kwargs):
        token = request.headers.get("Authorization") or ""
        # Lido a cada request (não no import) para não congelar o valor
        # quando o ambiente muda depois que o módulo já foi importado.
        token_esperado = os.environ.get("ADMIN_TOKEN") or ""
        if not token_esperado or not hmac.compare_digest(token, token_esperado):
            return jsonify(error="Acesso negado."), 403
        return func(*args, **kwargs)
    return wrapper

def _carregar_pedido_itens(pedido_id: int, loja_codigo: str):
    pedido = get_pedido(pedido_id) if callable(get_pedido) else None
    if not pedido or pedido.get('loja_codigo') != loja_codigo:
        return None, None
    itens = get_itens_pedido(pedido_id) if callable(get_itens_pedido) else []
    return pedido, itens

# -----------------------------------------------------------------------------
# Rotas de Autenticação e Servir Arquivos
# -----------------------------------------------------------------------------
@app.post("/auth/signup")
def auth_signup():
    data = request.get_json(silent=True) or {}
    nome_loja = (data.get("nome_loja") or "").strip()
    nome = (data.get("nome") or "").strip()
    login = (data.get("login") or "").strip()
    senha = (data.get("senha") or "").strip()
    if not nome_loja or not nome or not login or not senha:
        return jsonify(error="Nome da loja, seu nome, login e senha são obrigatórios."), 400
    senha_hash = generate_password_hash(senha)
    conn = get_conn()
    try:
        # Loja (tenant, pendente de aprovação) e primeiro usuário admin
        # nascem juntos, na mesma transação.
        conn.execute("BEGIN")
        conn.execute("INSERT INTO lojas (codigo, nome) VALUES (?, ?)", (nome_loja, nome_loja))
        conn.execute(
            "INSERT INTO usuarios (loja, nome, login, senha_hash, papel, taxa_comissao, ativo) VALUES (?, ?, ?, ?, 'admin', 0, 1)",
            (nome_loja, nome, login, senha_hash),
        )
        conn.commit()
    except sqlite3.IntegrityError as e:
        conn.rollback()
        if "usuarios" in str(e):
            return jsonify(error="Já existe um usuário com este login nesta loja."), 409
        return jsonify(error="Já existe uma loja com este nome."), 409
    finally:
        conn.close()
    return jsonify(ok=True, message="Loja criada com sucesso."), 201

@app.post("/auth/login")
@limiter.limit("5 per minute; 30 per hour")
def auth_login():
    data = request.get_json(silent=True) or {}
    nome_loja = (data.get("nome_loja") or "").strip()
    login = (data.get("login") or "").strip()
    senha = (data.get("senha") or "").strip()
    if not nome_loja or not login or not senha:
        return jsonify(error="Nome da loja, login e senha são obrigatórios."), 400
    loja_db = get_loja_by_codigo(nome_loja)
    usuario = get_usuario(nome_loja, login)
    if (
        loja_db and loja_db.get("status") == "aprovado"
        and usuario and usuario.get("ativo")
        # Operador com convite ainda não aceito não tem senha: não loga.
        and usuario.get("senha_hash")
        and check_password_hash(usuario["senha_hash"], senha)
    ):
        session["loja_codigo"] = loja_db["codigo"]
        session["usuario_id"] = usuario["id"]
        session["papel"] = usuario["papel"]
        return jsonify(ok=True, loja=loja_db["codigo"], usuario=usuario["nome"], papel=usuario["papel"])
    # Mensagem única de propósito: não revelar se falhou loja, usuário ou senha.
    return jsonify(error="Credenciais inválidas."), 401

@app.post("/auth/logout")
def auth_logout():
    session.clear()
    return jsonify(ok=True)

@app.get("/")
def index():
    if current_loja():
        return send_from_directory(HTML_DIR, "index.html")
    return redirect(url_for("acesso_page"))

@app.get("/acesso", endpoint="acesso_page")
def acesso_page():
    return send_from_directory(HTML_DIR, "acesso.html")

# -----------------------------------------------------------------------------
# Convite de primeiro acesso (público, protegido só pelo token)
# -----------------------------------------------------------------------------
CONVITE_VALIDADE = timedelta(days=7)
CONVITE_SENHA_MIN = 6


def _agora_utc_str():
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")


def _novo_convite():
    """(token, expira) para um convite novo."""
    expira = (datetime.now(timezone.utc) + CONVITE_VALIDADE).strftime("%Y-%m-%d %H:%M:%S")
    return secrets.token_urlsafe(32), expira


def _usuario_do_convite(conn, token):
    """Usuário dono do token, se o convite existe, não expirou e o usuário está ativo."""
    if not token:
        return None
    return conn.execute(
        "SELECT id, loja, nome, login FROM usuarios "
        "WHERE convite_token = ? AND convite_expira > ? AND ativo = 1",
        (token, _agora_utc_str()),
    ).fetchone()


@app.get("/convite/<token>")
def convite_page(token):
    conn = get_conn()
    try:
        usuario = _usuario_do_convite(conn, token)
    finally:
        conn.close()
    if not usuario:
        return render_template("convite.html", valido=False), 404
    return render_template(
        "convite.html", valido=True, nome=usuario["nome"],
        login=usuario["login"], loja=usuario["loja"],
    )


@app.post("/convite/<token>")
@limiter.limit("10 per minute")
def convite_definir_senha(token):
    data = request.get_json(silent=True) or {}
    senha = (data.get("senha") or "").strip()
    conn = get_conn()
    try:
        usuario = _usuario_do_convite(conn, token)
        if not usuario:
            return jsonify(error="Convite inválido ou expirado. Peça um novo link ao administrador."), 404
        if len(senha) < CONVITE_SENHA_MIN:
            return jsonify(error=f"A senha precisa ter pelo menos {CONVITE_SENHA_MIN} caracteres."), 400
        # O WHERE repete o token: se dois envios correrem juntos, só um grava.
        cur = conn.execute(
            "UPDATE usuarios SET senha_hash = ?, convite_token = NULL, convite_expira = NULL "
            "WHERE id = ? AND convite_token = ?",
            (generate_password_hash(senha), usuario["id"], token),
        )
        if cur.rowcount != 1:
            conn.rollback()
            return jsonify(error="Convite inválido ou expirado. Peça um novo link ao administrador."), 404
        conn.commit()
        return jsonify(ok=True, loja=usuario["loja"], login=usuario["login"],
                       message="Senha definida. Agora é só entrar.")
    finally:
        conn.close()


@app.get("/admin", endpoint="admin_page")
def admin_page():
    return send_from_directory(HTML_DIR, "admin.html")

@app.get("/<path:folder>/<path:filename>")
def serve_static_folder(folder, filename):
    if folder in ["css", "js", "html"]:
        return send_from_directory(FRONT_DIR / folder, filename)
    abort(404)

# -----------------------------------------------------------------------------
# API: Clientes, Vendedores
# -----------------------------------------------------------------------------
@app.route("/clientes", methods=["GET", "POST"])
@require_login
def clientes_api():
    loja = current_loja()
    conn = get_conn()
    if request.method == "GET":
        search_query = request.args.get("search", "").strip()
        sql = "SELECT id, nome, telefone, email FROM clientes WHERE loja = ? AND nome LIKE ? ORDER BY nome ASC"
        clientes = [dict(r) for r in conn.execute(sql, (loja, f"%{search_query}%")).fetchall()]
        conn.close()
        return jsonify(clientes)
    if request.method == "POST":
        data = request.get_json(silent=True) or {}
        nome = (data.get("nome") or "").strip()
        if not nome:
            conn.close(); return jsonify(error="Nome é obrigatório."), 400
        try:
            cur = conn.execute("INSERT INTO clientes (nome, telefone, email, loja) VALUES (?, ?, ?, ?)", (nome, data.get("telefone"), data.get("email"), loja))
            new_id = cur.lastrowid
            conn.commit()
            return jsonify(id=new_id, **data), 201
        except sqlite3.Error as e:
            conn.rollback(); return jsonify(error=f"Erro no banco de dados: {e}"), 500
        finally: conn.close()

@app.route("/clientes/<int:cliente_id>", methods=["DELETE"])
@require_login
def cliente_delete_api(cliente_id: int):
    loja = current_loja()
    conn = get_conn()
    try:
        cur = conn.execute("DELETE FROM clientes WHERE id = ? AND loja = ?", (cliente_id, loja))
        if cur.rowcount == 0:
            conn.close(); return jsonify(error="Cliente não encontrado."), 404
        conn.commit()
        return jsonify(ok=True, message="Cliente removido com sucesso.")
    except sqlite3.Error as e:
        conn.rollback(); return jsonify(error=f"Erro de banco de dados: {e}"), 500
    finally: conn.close()

@app.put("/clientes/<int:cliente_id>")
@require_login
def cliente_update_api(cliente_id: int):
    loja = current_loja()
    data = request.get_json(silent=True) or {}
    telefone = data.get("telefone", "")
    email = data.get("email", "")

    conn = get_conn()
    try:
        cur = conn.execute("UPDATE clientes SET telefone = ?, email = ? WHERE id = ? AND loja = ?",
                           (telefone, email, cliente_id, loja))
        if cur.rowcount == 0:
            conn.close(); return jsonify(error="Cliente não encontrado."), 404
        
        conn.commit()
        return jsonify(ok=True, message="Cliente atualizado com sucesso.")
    except sqlite3.Error as e:
        conn.rollback(); return jsonify(error=f"Erro de banco de dados: {e}"), 500
    finally:
        conn.close()

# -----------------------------------------------------------------------------
# API: Fiado (saldo devedor por cliente e pagamentos)
# -----------------------------------------------------------------------------
def _saldo_cliente(conn, loja, cliente_id):
    total_fiado = conn.execute(
        "SELECT COALESCE(SUM(total), 0) AS t FROM pedidos WHERE cliente_id = ? AND loja = ? AND pago = 0",
        (cliente_id, loja),
    ).fetchone()["t"]
    total_pago = conn.execute(
        "SELECT COALESCE(SUM(valor), 0) AS t FROM pagamentos WHERE cliente_id = ? AND loja = ?",
        (cliente_id, loja),
    ).fetchone()["t"]
    return round(total_fiado - total_pago, 2)


@app.post("/api/clientes/<int:cliente_id>/pagamentos")
@require_login
@require_perm("fiado_gerir")
def cliente_pagamento_criar(cliente_id: int):
    loja = current_loja()
    data = request.get_json(silent=True) or {}
    valor = _ptbr_to_float(data.get("valor"))
    data_pagamento = (data.get("data") or "").strip()
    if valor is None or valor <= 0:
        return jsonify(error="Valor do pagamento deve ser maior que zero."), 400
    if not data_pagamento:
        return jsonify(error="Data do pagamento é obrigatória."), 400

    conn = get_conn()
    try:
        if not conn.execute("SELECT 1 FROM clientes WHERE id = ? AND loja = ?", (cliente_id, loja)).fetchone():
            return jsonify(error="Cliente inválido para esta loja."), 404
        cur = conn.execute(
            "INSERT INTO pagamentos (loja, cliente_id, valor, data) VALUES (?, ?, ?, ?)",
            (loja, cliente_id, valor, data_pagamento),
        )
        conn.commit()
        saldo = _saldo_cliente(conn, loja, cliente_id)
        return jsonify(id=cur.lastrowid, cliente_id=cliente_id, valor=round(valor, 2),
                        data=data_pagamento, saldo=saldo), 201
    except sqlite3.Error as e:
        conn.rollback(); return jsonify(error=f"Erro de banco de dados: {e}"), 500
    finally:
        conn.close()


@app.get("/api/clientes/<int:cliente_id>/saldo")
@require_login
@require_perm("fiado_gerir")
def cliente_saldo_api(cliente_id: int):
    loja = current_loja()
    conn = get_conn()
    try:
        if not conn.execute("SELECT 1 FROM clientes WHERE id = ? AND loja = ?", (cliente_id, loja)).fetchone():
            return jsonify(error="Cliente inválido para esta loja."), 404
        saldo = _saldo_cliente(conn, loja, cliente_id)
    finally:
        conn.close()
    return jsonify(cliente_id=cliente_id, saldo=saldo)


@app.get("/api/fiado")
@require_login
@require_perm("fiado_gerir")
def fiado_listar():
    loja = current_loja()
    conn = get_conn()
    try:
        rows = conn.execute("""
            SELECT * FROM (
                SELECT c.id AS cliente_id, c.nome AS nome,
                       COALESCE(pf.total, 0) - COALESCE(pg.total, 0) AS saldo
                FROM clientes c
                LEFT JOIN (
                    SELECT cliente_id, SUM(total) AS total FROM pedidos
                    WHERE loja = ? AND pago = 0 GROUP BY cliente_id
                ) pf ON pf.cliente_id = c.id
                LEFT JOIN (
                    SELECT cliente_id, SUM(valor) AS total FROM pagamentos
                    WHERE loja = ? GROUP BY cliente_id
                ) pg ON pg.cliente_id = c.id
                WHERE c.loja = ?
            )
            WHERE saldo > 0
            ORDER BY saldo DESC
        """, (loja, loja, loja)).fetchall()
    finally:
        conn.close()
    devedores = [{"cliente_id": r["cliente_id"], "nome": r["nome"], "saldo": round(r["saldo"], 2)} for r in rows]
    total = round(sum(d["saldo"] for d in devedores), 2)
    return jsonify(devedores=devedores, total=total)


@app.post("/pedidos/<int:pedido_id>/quitar")
@require_login
@require_perm("fiado_gerir")
def pedido_quitar(pedido_id: int):
    loja = current_loja()
    conn = get_conn()
    try:
        cur = conn.execute("UPDATE pedidos SET pago = 1 WHERE id = ? AND loja = ?", (pedido_id, loja))
        if cur.rowcount == 0:
            return jsonify(error="Pedido não encontrado."), 404
        conn.commit()
        return jsonify(ok=True, message="Pedido marcado como pago.")
    except sqlite3.Error as e:
        conn.rollback(); return jsonify(error=f"Erro de banco de dados: {e}"), 500
    finally:
        conn.close()


@app.get("/me")
@require_login
def me_api():
    usuario = get_usuario_by_id(session.get("usuario_id"))
    return jsonify(
        loja=current_loja(),
        usuario_id=usuario["id"],
        nome=usuario["nome"],
        papel=usuario["papel"],
        # O menu do front é montado a partir daqui: cada tela pede uma
        # capacidade, e o que o papel não tem simplesmente não aparece.
        permissoes=permissoes_do_papel(usuario["papel"]),
    )

# -----------------------------------------------------------------------------
# API: Usuários (gestão de operadores — somente admin da loja)
# -----------------------------------------------------------------------------
@app.route("/usuarios", methods=["GET", "POST"])
@require_login
@require_perm("usuarios_gerir")
def usuarios_api():
    loja = current_loja()
    conn = get_conn()
    if request.method == "GET":
        usuarios = [dict(r) for r in conn.execute(
            "SELECT id, nome, login, papel, taxa_comissao, ativo, "
            "(senha_hash IS NULL) AS convite_pendente "
            "FROM usuarios WHERE loja = ? ORDER BY nome ASC",
            (loja,)
        ).fetchall()]
        conn.close(); return jsonify(usuarios)
    if request.method == "POST":
        data = request.get_json(silent=True) or {}
        nome = (data.get("nome") or "").strip()
        login = (data.get("login") or "").strip()
        taxa = _ptbr_to_float(data.get("taxa_comissao"))
        if taxa is None: taxa = 0.0
        if not nome or not login:
            conn.close(); return jsonify(error="Nome e login são obrigatórios."), 400
        if taxa < 0:
            conn.close(); return jsonify(error="Taxa de comissão não pode ser negativa."), 400
        # O operador nasce SEM senha: quem a define é ele, pelo link do convite.
        token, expira = _novo_convite()
        try:
            cur = conn.execute(
                "INSERT INTO usuarios (loja, nome, login, senha_hash, papel, taxa_comissao, ativo, convite_token, convite_expira) "
                "VALUES (?, ?, ?, NULL, 'operador', ?, 1, ?, ?)",
                (loja, nome, login, taxa, token, expira),
            )
            new_id = cur.lastrowid
            conn.commit()
            convite_path = f"/convite/{token}"
            return jsonify(
                id=new_id, nome=nome, login=login, papel="operador", taxa_comissao=taxa,
                convite_token=token, convite_path=convite_path,
                convite_url=request.host_url.rstrip("/") + convite_path,
                convite_expira=expira,
            ), 201
        except sqlite3.IntegrityError:
            conn.rollback(); return jsonify(error=f"Já existe um usuário com o login '{login}' nesta loja."), 409
        finally: conn.close()

@app.put("/usuarios/<int:usuario_id>")
@require_login
@require_perm("usuarios_gerir")
def usuario_update_api(usuario_id: int):
    loja = current_loja()
    data = request.get_json(silent=True) or {}
    conn = get_conn()
    try:
        alvo = conn.execute(
            "SELECT id, nome, papel, taxa_comissao, ativo FROM usuarios WHERE id = ? AND loja = ?",
            (usuario_id, loja)
        ).fetchone()
        if not alvo:
            return jsonify(error="Usuário não encontrado."), 404

        nome = (data.get("nome") or "").strip() or alvo["nome"]
        taxa = _ptbr_to_float(data.get("taxa_comissao"))
        if taxa is None: taxa = alvo["taxa_comissao"]
        if taxa < 0:
            return jsonify(error="Taxa de comissão não pode ser negativa."), 400
        # O admin é o dono: a comissão é despesa dele, não receita.
        # A taxa do admin fica travada em 0 e não é editável.
        if alvo["papel"] == "admin":
            if taxa:
                return jsonify(error="O admin é o dono da loja e não recebe comissão."), 400
            taxa = 0.0
        ativo = 1 if data.get("ativo", alvo["ativo"]) in (1, True, "1", "true") else 0
        if usuario_id == session.get("usuario_id") and not ativo:
            return jsonify(error="Você não pode desativar a si mesmo."), 400

        conn.execute(
            "UPDATE usuarios SET nome = ?, taxa_comissao = ?, ativo = ? WHERE id = ? AND loja = ?",
            (nome, taxa, ativo, usuario_id, loja),
        )
        senha_nova = (data.get("senha") or "").strip()
        if senha_nova:
            # Senha definida pelo admin substitui um convite pendente.
            conn.execute(
                "UPDATE usuarios SET senha_hash = ?, convite_token = NULL, convite_expira = NULL "
                "WHERE id = ? AND loja = ?",
                (generate_password_hash(senha_nova), usuario_id, loja),
            )
        conn.commit()
        return jsonify(ok=True, message="Usuário atualizado.")
    except sqlite3.Error as e:
        conn.rollback(); return jsonify(error=f"Erro de banco de dados: {e}"), 500
    finally:
        conn.close()

# -----------------------------------------------------------------------------
# API: Pedidos com Lógica de Estoque
# -----------------------------------------------------------------------------
@app.route("/pedidos", methods=["GET", "POST"])
@require_login
def pedidos_api():
    loja = current_loja()
    conn = get_conn()
    if request.method == "GET":
        sort_by = request.args.get("sort", "data_desc")
        order_options = {"data_desc": "p.id DESC", "data_asc": "p.id ASC", "cliente_asc": "cliente_nome ASC, p.id DESC", "total_desc": "p.total DESC", "total_asc": "p.total ASC"}
        order_clause = order_options.get(sort_by, "p.id DESC")
        # ?status=orcamento lista os orçamentos; o default são os pedidos.
        status = "orcamento" if request.args.get("status") == "orcamento" else "pedido"
        sql = f"SELECT p.id, p.data_iso, p.tecido, p.total, p.desconto, p.comissao_valor, p.status, c.nome AS cliente_nome, COALESCE(u.nome, '') AS vendedor_nome FROM pedidos p JOIN clientes c ON c.id = p.cliente_id LEFT JOIN usuarios u ON u.id = p.usuario_id WHERE p.loja = ? AND p.status = ? ORDER BY {order_clause} LIMIT 200"
        pedidos_raw = conn.execute(sql, (loja, status)).fetchall()
        pedidos = [dict(p) for p in pedidos_raw]
        conn.close()
        return jsonify(pedidos)
    
    if request.method == "POST":
        if not pode("vender"):
            conn.close()
            return jsonify(error="Acesso restrito: permissão insuficiente."), 403
        data = request.get_json(silent=True) or {}
        cliente_id = data.get("cliente_id")
        preco_unitario = _ptbr_to_float(data.get("preco_unitario"))
        itens_in = data.get("itens") or []
        descontar_estoque = data.get("descontar_estoque", False)
        # Orçamento: não baixa estoque e não congela comissão (fica 0 até a
        # conversão em POST /pedidos/<id>/converter).
        is_orcamento = data.get("is_orcamento") in (True, 1, "1", "true")
        tecido_nome = (data.get("tecido") or "").strip()
        # Fiado é só forma de pagamento: não muda baixa de estoque nem comissão.
        pago = 0 if data.get("pago", True) in (0, False, "0", "false", "False") else 1

        if not all([cliente_id, preco_unitario is not None, itens_in, tecido_nome]):
            conn.close()
            return jsonify(error="Dados incompletos."), 400

        if not conn.execute("SELECT 1 FROM clientes WHERE id = ? AND loja = ?", (cliente_id, loja)).fetchone():
            conn.close()
            return jsonify(error="Cliente inválido para esta loja."), 400

        total_kg = sum(_ptbr_to_float(it.get("peso", 0)) for it in itens_in)
        desconto = _ptbr_to_float(data.get("desconto")) or 0.0
        total = round((total_kg * preco_unitario) - desconto, 2)

        # O vendedor é sempre o usuário logado; a comissão é congelada aqui
        # com a taxa vigente — mudanças futuras não afetam este pedido.
        usuario_id = session.get("usuario_id")
        if is_orcamento:
            comissao_taxa = comissao_valor = 0.0
        else:
            comissao_taxa = _taxa_comissao(get_usuario_by_id(usuario_id))
            comissao_valor = round(total * comissao_taxa / 100, 2)
        status = "orcamento" if is_orcamento else "pedido"

        try:
            conn.execute("BEGIN")

            if descontar_estoque and not is_orcamento:
                _baixar_estoque(conn, loja, tecido_nome,
                                [(it.get("cor"), _ptbr_to_float(it.get("peso", 0))) for it in itens_in])

            now_iso = datetime.now().strftime("%Y-%m-%d %H:%M")
            cur_pedido = conn.execute(
                "INSERT INTO pedidos (cliente_id, tecido, quantidade, preco_unitario, total, desconto, loja, data_iso, usuario_id, comissao_taxa, comissao_valor, pago, status) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (cliente_id, tecido_nome, len(itens_in), preco_unitario, total, desconto, loja, now_iso, usuario_id, comissao_taxa, comissao_valor, pago, status)
            )
            pedido_id = cur_pedido.lastrowid

            itens_to_insert = [(pedido_id, tecido_nome, it.get("cor"), _ptbr_to_float(it.get("peso"))) for it in itens_in]
            conn.executemany("INSERT INTO itens_pedido (pedido_id, descricao, cor, peso_kg) VALUES (?, ?, ?, ?)", itens_to_insert)
            
            conn.commit()
        except EstoqueInsuficiente as e:
            conn.rollback()
            return jsonify(error=str(e)), 409
        except sqlite3.Error as e:
            conn.rollback()
            return jsonify(error=f"Erro de banco de dados: {e}"), 500
        finally:
            conn.close()

        return jsonify(id=pedido_id, total=total, status=status), 201

@app.post("/pedidos/<int:pedido_id>/converter")
@require_login
@require_perm("vender")
def pedido_converter_api(pedido_id: int):
    """Orçamento -> pedido: baixa o estoque (mesma validação por peso) e
    congela a comissão pela taxa ATUAL de quem fez o orçamento."""
    loja = current_loja()
    conn = get_conn()
    try:
        # IMMEDIATE: trava a escrita já na leitura, para duas conversões
        # simultâneas não baixarem o estoque duas vezes.
        conn.execute("BEGIN IMMEDIATE")
        pedido = conn.execute(
            "SELECT id, tecido, total, usuario_id, status FROM pedidos WHERE id = ? AND loja = ?",
            (pedido_id, loja),
        ).fetchone()
        if not pedido:
            conn.rollback()
            return jsonify(error="Orçamento não encontrado."), 404
        if pedido["status"] != "orcamento":
            conn.rollback()
            return jsonify(error="Este pedido já foi convertido."), 409

        itens = conn.execute("SELECT cor, peso_kg FROM itens_pedido WHERE pedido_id = ?", (pedido_id,)).fetchall()
        _baixar_estoque(conn, loja, pedido["tecido"], [(i["cor"], i["peso_kg"]) for i in itens])

        vendedor = get_usuario_by_id(pedido["usuario_id"]) if pedido["usuario_id"] else None
        comissao_taxa = _taxa_comissao(vendedor)
        comissao_valor = round((pedido["total"] or 0) * comissao_taxa / 100, 2)
        # A venda acontece na conversão: data_iso passa a ser agora, para o
        # pedido cair no período certo dos relatórios.
        now_iso = datetime.now().strftime("%Y-%m-%d %H:%M")
        cur = conn.execute(
            "UPDATE pedidos SET status = 'pedido', comissao_taxa = ?, comissao_valor = ?, data_iso = ? "
            "WHERE id = ? AND loja = ? AND status = 'orcamento'",
            (comissao_taxa, comissao_valor, now_iso, pedido_id, loja),
        )
        if cur.rowcount != 1:
            raise EstoqueInsuficiente("Este pedido já foi convertido.")
        conn.commit()
    except EstoqueInsuficiente as e:
        conn.rollback()
        return jsonify(error=str(e)), 409
    except sqlite3.Error as e:
        conn.rollback()
        return jsonify(error=f"Erro de banco de dados: {e}"), 500
    finally:
        conn.close()
    return jsonify(ok=True, id=pedido_id, status="pedido", comissao_valor=comissao_valor)

@app.route("/pedidos/<int:pedido_id>", methods=["GET", "PUT", "DELETE"])
@require_login
def pedido_single_api(pedido_id: int):
    loja = current_loja()
    conn = get_conn()
    if request.method == "GET":
        pedido, itens = _carregar_pedido_itens(pedido_id, loja)
        if not pedido:
            conn.close(); return jsonify(error="Pedido não encontrado."), 404
        conn.close(); return jsonify(pedido=dict(pedido), itens=itens)
        
    if request.method == "DELETE":
        try:
            cur = conn.execute("DELETE FROM pedidos WHERE id = ? AND loja = ?", (pedido_id, loja))
            if cur.rowcount == 0:
                conn.close(); return jsonify(error="Pedido não encontrado."), 404
            conn.commit()
        except sqlite3.Error as e:
            conn.rollback(); return jsonify(error=f"Erro no banco de dados: {e}"), 500
        finally: conn.close()
        return jsonify(ok=True)
    
    if request.method == "PUT":
        data = request.get_json(silent=True) or {}
        # NOTA: A lógica aqui NÃO afeta o estoque. É apenas para correção de dados do romaneio.
        try:
            cliente_id = data.get("cliente_id")
            tecido_nome = (data.get("tecido") or "").strip()
            preco_unitario = _ptbr_to_float(data.get("preco_unitario"))
            desconto = _ptbr_to_float(data.get("desconto")) or 0.0
            itens_in = data.get("itens") or []

            total_kg = sum(_ptbr_to_float(it.get("peso", 0)) for it in itens_in)
            total = round((total_kg * preco_unitario) - desconto, 2)

            if not conn.execute("SELECT 1 FROM clientes WHERE id = ? AND loja = ?", (cliente_id, loja)).fetchone():
                return jsonify(error="Cliente inválido para esta loja."), 400

            conn.execute("BEGIN")
            # O vendedor (usuario_id) não muda na edição; a comissão é
            # recalculada sobre o novo total usando a taxa JÁ CONGELADA
            # no pedido, não a taxa atual do usuário.
            cur = conn.execute("""
                UPDATE pedidos SET cliente_id=?, tecido=?, preco_unitario=?, desconto=?, total=?,
                    comissao_valor = ROUND(? * comissao_taxa / 100, 2)
                WHERE id=? AND loja=?
            """, (cliente_id, tecido_nome, preco_unitario, desconto, total, total, pedido_id, loja))
            if cur.rowcount == 0:
                conn.rollback()
                return jsonify(error="Pedido não encontrado."), 404

            conn.execute("DELETE FROM itens_pedido WHERE pedido_id=?", (pedido_id,))
            
            itens_to_insert = [(pedido_id, tecido_nome, it.get("cor"), _ptbr_to_float(it.get("peso"))) for it in itens_in]
            conn.executemany("INSERT INTO itens_pedido (pedido_id, descricao, cor, peso_kg) VALUES (?, ?, ?, ?)", itens_to_insert)
            
            conn.commit()
            return jsonify(ok=True, message="Pedido atualizado.")
        except sqlite3.Error as e:
            conn.rollback()
            return jsonify(error=f"Erro no banco de dados: {e}"), 500
        finally:
            conn.close()

# -----------------------------------------------------------------------------
# API: Estoque
# -----------------------------------------------------------------------------
@app.get("/api/estoque")
@require_login
def get_estoque():
    loja = current_loja()
    conn = get_conn()
    tecidos_raw = conn.execute("SELECT id, nome_tecido FROM estoque_tecidos WHERE loja = ? ORDER BY nome_tecido", (loja,)).fetchall()
    
    estoque = []
    for tecido in tecidos_raw:
        cores = conn.execute(
            "SELECT id, nome_cor, peso_kg, qtd_pecas, estoque_minimo FROM estoque_cores WHERE tecido_id = ? ORDER BY nome_cor",
            (tecido['id'],)
        ).fetchall()
        estoque.append({
            "id": tecido['id'],
            "nome_tecido": tecido['nome_tecido'],
            # Alerta só quando há mínimo definido (> 0) e o peso chegou nele.
            "cores": [
                {**dict(c), "abaixo_minimo": c['estoque_minimo'] > 0 and c['peso_kg'] <= c['estoque_minimo']}
                for c in cores
            ]
        })
    conn.close()
    return jsonify(estoque)

@app.post("/api/estoque/tecidos")
@require_login
@require_perm("estoque_mutar")
def add_tecido_estoque():
    loja = current_loja()
    data = request.get_json()
    nome_tecido = (data.get("nome_tecido") or "").strip()
    if not nome_tecido:
        return jsonify(error="Nome do tecido é obrigatório."), 400
    
    conn = get_conn()
    try:
        cur = conn.execute("INSERT INTO estoque_tecidos (nome_tecido, loja) VALUES (?, ?)", (nome_tecido, loja))
        new_id = cur.lastrowid
        conn.commit()
        return jsonify(id=new_id, nome_tecido=nome_tecido), 201
    except sqlite3.IntegrityError:
        conn.rollback()
        return jsonify(error=f"O tecido '{nome_tecido}' já está cadastrado."), 409
    finally:
        conn.close()

@app.post("/api/estoque/cores")
@require_login
@require_perm("estoque_mutar")
def add_cor_estoque():
    loja = current_loja()
    data = request.get_json()
    tecido_id = data.get("tecido_id")
    nome_cor = (data.get("nome_cor") or "").strip()
    peso_kg = _ptbr_to_float(data.get("peso_kg")) or 0
    qtd_pecas = data.get("qtd_pecas") or 0
    
    if not all([tecido_id, nome_cor]):
        return jsonify(error="ID do tecido e nome da cor são obrigatórios."), 400

    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute("SELECT 1 FROM estoque_tecidos WHERE id = ? AND loja = ?", (tecido_id, loja))
        if not cur.fetchone():
            return jsonify(error="Tecido não encontrado nesta loja."), 404
        cur.execute("SELECT id, peso_kg, qtd_pecas FROM estoque_cores WHERE tecido_id = ? AND nome_cor = ?", (tecido_id, nome_cor))
        existing = cur.fetchone()
        
        if existing:
            novo_peso = existing['peso_kg'] + peso_kg
            novas_pecas = existing['qtd_pecas'] + qtd_pecas
            cur.execute("UPDATE estoque_cores SET peso_kg = ?, qtd_pecas = ? WHERE id = ?", (novo_peso, novas_pecas, existing['id']))
        else:
            cur.execute("INSERT INTO estoque_cores (tecido_id, nome_cor, peso_kg, qtd_pecas) VALUES (?, ?, ?, ?)",
                        (tecido_id, nome_cor, peso_kg, qtd_pecas))
        
        conn.commit()
        return jsonify(ok=True, message="Estoque atualizado."), 200
    except sqlite3.Error as e:
        conn.rollback()
        return jsonify(error=f"Erro de banco de dados: {e}"), 500
    finally:
        conn.close()

@app.delete("/api/estoque/tecidos/<int:tecido_id>")
@require_login
@require_perm("estoque_mutar")
def delete_tecido(tecido_id):
    loja = current_loja()
    conn = get_conn()
    cur = conn.execute("DELETE FROM estoque_tecidos WHERE id = ? AND loja = ?", (tecido_id, loja))
    if cur.rowcount == 0:
        conn.close()
        return jsonify(error="Tecido não encontrado ou não pertence a esta loja."), 404
    conn.commit()
    conn.close()
    return jsonify(ok=True)

@app.delete("/api/estoque/cores/<int:cor_id>")
@require_login
@require_perm("estoque_mutar")
def delete_cor(cor_id):
    loja = current_loja()
    conn = get_conn()
    cur = conn.execute(
        "DELETE FROM estoque_cores WHERE id = ? AND tecido_id IN (SELECT id FROM estoque_tecidos WHERE loja = ?)",
        (cor_id, loja)
    )
    if cur.rowcount == 0:
        conn.close()
        return jsonify(error="Cor não encontrada."), 404
    conn.commit()
    conn.close()
    return jsonify(ok=True)

@app.put("/api/estoque/cores/<int:cor_id>")
@require_login
@require_perm("estoque_mutar")
def update_cor_estoque(cor_id):
    loja = current_loja()
    data = request.get_json()
    novo_peso_kg = _ptbr_to_float(data.get("peso_kg"))
    novas_qtd_pecas = data.get("qtd_pecas")

    if novo_peso_kg is None or novas_qtd_pecas is None or novo_peso_kg < 0 or novas_qtd_pecas < 0:
        return jsonify(error="Valores de peso e quantidade de peças são obrigatórios e não podem ser negativos."), 400

    conn = get_conn()
    try:
        cur = conn.execute(
            "UPDATE estoque_cores SET peso_kg = ?, qtd_pecas = ? WHERE id = ? AND tecido_id IN (SELECT id FROM estoque_tecidos WHERE loja = ?)",
            (novo_peso_kg, novas_qtd_pecas, cor_id, loja)
        )
        
        if cur.rowcount == 0:
            conn.close(); return jsonify(error="Cor não encontrada no estoque."), 404
            
        conn.commit()
        return jsonify(ok=True, message="Estoque ajustado com sucesso.")
    except sqlite3.Error as e:
        conn.rollback(); return jsonify(error=f"Erro de banco de dados: {e}"), 500
    finally:
        conn.close()

@app.put("/api/estoque/cores/<int:cor_id>/minimo")
@require_login
@require_perm("estoque_mutar")
def update_cor_minimo(cor_id):
    loja = current_loja()
    data = request.get_json(silent=True) or {}
    minimo = _peso_to_float(data.get("estoque_minimo"))
    if minimo is None or minimo < 0:
        return jsonify(error="Estoque mínimo é obrigatório e não pode ser negativo (0 desliga o alerta)."), 400

    conn = get_conn()
    try:
        cur = conn.execute(
            "UPDATE estoque_cores SET estoque_minimo = ? WHERE id = ? AND tecido_id IN (SELECT id FROM estoque_tecidos WHERE loja = ?)",
            (minimo, cor_id, loja)
        )
        if cur.rowcount == 0:
            return jsonify(error="Cor não encontrada no estoque."), 404
        conn.commit()
        return jsonify(ok=True, estoque_minimo=minimo)
    except sqlite3.Error as e:
        conn.rollback(); return jsonify(error=f"Erro de banco de dados: {e}"), 500
    finally:
        conn.close()

# -----------------------------------------------------------------------------
# API: Entrada de mercadoria (recebimento avulso)
# -----------------------------------------------------------------------------
# Duas formas de lançar, uma só confirmação: o lote digitado vai direto para
# POST /api/estoque/entradas; a planilha passa antes por .../importar, que só
# lê e devolve as linhas para o gerente/admin conferir — nada é gravado ali.
ENTRADA_MAX_LINHAS = 2000
PLANILHA_MAX_BYTES = 2 * 1024 * 1024

# Cabeçalho normalizado (minúsculas, sem acento/espaço) -> campo da linha.
_COLUNAS_PLANILHA = {
    "tecido": "tecido",
    "cor": "cor",
    "peso": "peso", "pesokg": "peso", "kg": "peso",
    "idrolo": "id_rolo", "iddorolo": "id_rolo", "rolo": "id_rolo", "rollid": "id_rolo",
}

def _peso_to_float(val):
    """Peso vindo de digitação ou planilha: aceita '12,5', '1.234,5' e '12.5'.

    _ptbr_to_float trata todo ponto como milhar ('12.5' -> 125), o que é
    perigoso para CSV exportado com ponto decimal. Aqui, sem vírgula, o ponto
    é decimal.
    """
    if val is None: return None
    if isinstance(val, bool): return None
    if isinstance(val, (int, float)): return float(val)
    s = str(val).strip()
    if "," in s: return _ptbr_to_float(s)
    try: return float(re.sub(r"[^\d.\-]", "", s))
    except ValueError: return None

def _normalizar_cabecalho(nome):
    s = unicodedata.normalize("NFKD", str(nome or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", s.lower())

def _texto_celula(val):
    """Célula -> texto. Números inteiros do Excel (ex.: id do rolo 123.0) viram '123'."""
    if val is None: return ""
    if isinstance(val, float) and val.is_integer(): return str(int(val))
    return str(val).strip()

def _validar_linhas_entrada(linhas_in):
    """Normaliza as linhas e devolve (linhas, erros). erros = [{linha, erro}], 1-based."""
    linhas, erros = [], []
    for i, it in enumerate(linhas_in, start=1):
        if not isinstance(it, dict):
            erros.append({"linha": i, "erro": "Linha inválida."}); continue
        tecido = _texto_celula(it.get("tecido"))
        cor = _texto_celula(it.get("cor"))
        peso = _peso_to_float(it.get("peso"))
        id_rolo = _texto_celula(it.get("id_rolo")) or None
        faltando = [n for n, v in (("tecido", tecido), ("cor", cor)) if not v]
        if faltando:
            erros.append({"linha": i, "erro": f"Campo obrigatório vazio: {', '.join(faltando)}."}); continue
        if peso is None or peso <= 0:
            erros.append({"linha": i, "erro": f"Peso inválido: '{_texto_celula(it.get('peso'))}'."}); continue
        linhas.append({"tecido": tecido, "cor": cor, "peso": round(peso, 3), "id_rolo": id_rolo})
    return linhas, erros

def _ler_planilha(nome_arquivo, conteudo):
    """Lê csv/xlsx e devolve uma lista de dicts crus {tecido, cor, peso, id_rolo}."""
    ext = Path(nome_arquivo or "").suffix.lower()
    if ext == ".xlsx":
        try:
            wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        except Exception:
            raise ValueError("Não foi possível ler o arquivo .xlsx.")
        try:
            rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        finally:
            wb.close()
    elif ext == ".csv":
        try:
            texto = conteudo.decode("utf-8-sig")
        except UnicodeDecodeError:
            texto = conteudo.decode("latin-1")
        try:
            # Excel em pt-BR exporta com ';'; outros com ',' ou tab.
            dialeto = csv.Sniffer().sniff(texto[:4096], delimiters=";,\t")
        except csv.Error:
            dialeto = csv.excel
        rows = list(csv.reader(io.StringIO(texto), dialeto))
    else:
        raise ValueError("Formato não suportado: envie .csv ou .xlsx.")

    rows = [r for r in rows if any(_texto_celula(c) for c in r)]
    if not rows:
        raise ValueError("A planilha está vazia.")
    campos = [_COLUNAS_PLANILHA.get(_normalizar_cabecalho(c)) for c in rows[0]]
    faltando = [c for c in ("tecido", "cor", "peso") if c not in campos]
    if faltando:
        raise ValueError(f"Coluna(s) obrigatória(s) ausente(s) no cabeçalho: {', '.join(faltando)}.")
    return [
        {campo: valor for campo, valor in zip(campos, r) if campo}
        for r in rows[1:]
    ]

@app.post("/api/estoque/entradas/importar")
@require_login
@require_perm("estoque_mutar")
def importar_planilha_entrada():
    arquivo = request.files.get("arquivo")
    if not arquivo or not arquivo.filename:
        return jsonify(error="Envie o arquivo da planilha (campo 'arquivo')."), 400
    conteudo = arquivo.read(PLANILHA_MAX_BYTES + 1)
    if len(conteudo) > PLANILHA_MAX_BYTES:
        return jsonify(error="Planilha muito grande (máx. 2 MB)."), 400
    try:
        brutas = _ler_planilha(arquivo.filename, conteudo)
    except ValueError as e:
        return jsonify(error=str(e)), 400
    if len(brutas) > ENTRADA_MAX_LINHAS:
        return jsonify(error=f"Máximo de {ENTRADA_MAX_LINHAS} linhas por entrada."), 400
    linhas, erros = _validar_linhas_entrada(brutas)
    # A linha 1 da planilha é o cabeçalho: ajusta a numeração para o que o
    # usuário vê no Excel.
    for e in erros: e["linha"] += 1
    return jsonify(linhas=linhas, erros=erros)

@app.post("/api/estoque/entradas")
@require_login
@require_perm("estoque_mutar")
def registrar_entrada():
    loja = current_loja()
    data = request.get_json(silent=True) or {}
    fornecedor = (data.get("fornecedor") or "").strip() or None
    data_entrada = (data.get("data") or "").strip() or datetime.now().strftime("%Y-%m-%d")
    linhas_in = data.get("linhas")
    if not isinstance(linhas_in, list) or not linhas_in:
        return jsonify(error="Informe ao menos uma linha (tecido, cor, peso)."), 400
    if len(linhas_in) > ENTRADA_MAX_LINHAS:
        return jsonify(error=f"Máximo de {ENTRADA_MAX_LINHAS} linhas por entrada."), 400

    # Tudo ou nada: qualquer linha inválida recusa a entrada inteira.
    linhas, erros = _validar_linhas_entrada(linhas_in)
    if erros:
        return jsonify(error="Há linhas inválidas na entrada.", erros=erros), 400

    ids_rolo = [l["id_rolo"] for l in linhas if l["id_rolo"]]
    repetidos = sorted({r for r in ids_rolo if ids_rolo.count(r) > 1})
    if repetidos:
        return jsonify(error=f"Id de rolo repetido na entrada: {', '.join(repetidos)}."), 400

    conn = get_conn()
    try:
        conn.execute("BEGIN")
        if ids_rolo:
            marcas = ",".join("?" * len(ids_rolo))
            ja = [r["roll_ext_id"] for r in conn.execute(
                f"SELECT roll_ext_id FROM rolos WHERE loja = ? AND roll_ext_id IN ({marcas})",
                (loja, *ids_rolo),
            ).fetchall()]
            if ja:
                conn.rollback()
                return jsonify(error=f"Rolo(s) já recebido(s) antes: {', '.join(sorted(ja))}."), 409

        # A cor é casada sem diferenciar maiúsculas ('azul' credita 'Azul'),
        # para a planilha não criar cores duplicadas. O tecido precisa existir:
        # um nome de tecido errado vira erro, não um tecido novo silencioso.
        tecidos = {
            r["nome_tecido"].casefold(): (r["id"], r["nome_tecido"])
            for r in conn.execute("SELECT id, nome_tecido FROM estoque_tecidos WHERE loja = ?", (loja,)).fetchall()
        }
        desconhecidos = sorted({l["tecido"] for l in linhas if l["tecido"].casefold() not in tecidos})
        if desconhecidos:
            conn.rollback()
            return jsonify(error=f"Tecido(s) não cadastrado(s) no estoque: {', '.join(desconhecidos)}. Cadastre o tecido antes da entrada."), 400

        entrada_id = conn.execute(
            "INSERT INTO entradas (loja, fornecedor, data) VALUES (?, ?, ?)",
            (loja, fornecedor, data_entrada),
        ).lastrowid

        cores_cache = {}
        cores_criadas = []
        for l in linhas:
            tecido_id, tecido_nome = tecidos[l["tecido"].casefold()]
            if tecido_id not in cores_cache:
                cores_cache[tecido_id] = {
                    r["nome_cor"].casefold(): (r["id"], r["nome_cor"])
                    for r in conn.execute("SELECT id, nome_cor FROM estoque_cores WHERE tecido_id = ?", (tecido_id,)).fetchall()
                }
            cores = cores_cache[tecido_id]
            chave = l["cor"].casefold()
            if chave in cores:
                cor_id, cor_nome = cores[chave]
                conn.execute("UPDATE estoque_cores SET peso_kg = ROUND(peso_kg + ?, 3) WHERE id = ?", (l["peso"], cor_id))
            else:
                cor_nome = l["cor"]
                cor_id = conn.execute(
                    "INSERT INTO estoque_cores (tecido_id, nome_cor, peso_kg, qtd_pecas) VALUES (?, ?, ?, 0)",
                    (tecido_id, cor_nome, l["peso"]),
                ).lastrowid
                cores[chave] = (cor_id, cor_nome)
                cores_criadas.append({"tecido": tecido_nome, "cor": cor_nome})
            if l["id_rolo"]:
                conn.execute(
                    "INSERT INTO rolos (loja, entrada_id, tecido, cor, roll_ext_id, peso) VALUES (?, ?, ?, ?, ?, ?)",
                    (loja, entrada_id, tecido_nome, cor_nome, l["id_rolo"], l["peso"]),
                )
        conn.commit()
    except sqlite3.Error as e:
        conn.rollback()
        return jsonify(error=f"Erro de banco de dados: {e}"), 500
    finally:
        conn.close()

    return jsonify(
        id=entrada_id,
        linhas=len(linhas),
        peso_total=round(sum(l["peso"] for l in linhas), 3),
        rolos=len(ids_rolo),
        cores_criadas=cores_criadas,
    ), 201

# -----------------------------------------------------------------------------
# API: Encomendas (o que está PREVISTO chegar)
# -----------------------------------------------------------------------------
# O previsto nunca mexe no estoque. Ao "receber", os pesos REAIS (que podem
# diferir do previsto) creditam o estoque pelo MESMO caminho da entrada
# (soma em estoque_cores, cria cor se preciso, grava rolos se vier id) — o
# recebimento inclusive grava uma linha em `entradas`, então ele aparece no
# histórico de entradas como qualquer outro recebimento.
def _get_encomenda(encomenda_id, loja):
    conn = get_conn()
    try:
        enc = conn.execute(
            "SELECT id, loja, fornecedor, data_prevista, status, created_at FROM encomendas WHERE id = ? AND loja = ?",
            (encomenda_id, loja),
        ).fetchone()
        if not enc:
            return None, None
        itens = conn.execute(
            "SELECT id, tecido, cor, peso_previsto, rolos_previstos FROM encomenda_itens WHERE encomenda_id = ? ORDER BY id",
            (encomenda_id,),
        ).fetchall()
        return dict(enc), [dict(i) for i in itens]
    finally:
        conn.close()

def _validar_itens_encomenda(itens_in):
    """Normaliza os itens previstos e devolve (itens, erros). erros = [{linha, erro}], 1-based."""
    itens, erros = [], []
    for i, it in enumerate(itens_in, start=1):
        if not isinstance(it, dict):
            erros.append({"linha": i, "erro": "Item inválido."}); continue
        tecido = _texto_celula(it.get("tecido"))
        cor = _texto_celula(it.get("cor"))
        peso_previsto = _peso_to_float(it.get("peso_previsto"))
        rolos_previstos_raw = it.get("rolos_previstos")
        faltando = [n for n, v in (("tecido", tecido), ("cor", cor)) if not v]
        if faltando:
            erros.append({"linha": i, "erro": f"Campo obrigatório vazio: {', '.join(faltando)}."}); continue
        if peso_previsto is None or peso_previsto <= 0:
            erros.append({"linha": i, "erro": f"Peso previsto inválido: '{_texto_celula(it.get('peso_previsto'))}'."}); continue
        try:
            rolos_previstos = int(rolos_previstos_raw) if rolos_previstos_raw not in (None, "") else None
            if rolos_previstos is not None and rolos_previstos < 0:
                raise ValueError
        except (TypeError, ValueError):
            erros.append({"linha": i, "erro": "Rolos previstos inválido."}); continue
        itens.append({"tecido": tecido, "cor": cor, "peso_previsto": round(peso_previsto, 3), "rolos_previstos": rolos_previstos})
    return itens, erros

def _validar_linhas_recebimento(linhas_in, itens_validos):
    """Como _validar_linhas_entrada, mas cada linha pode trazer o item_id previsto que credita."""
    linhas, erros = [], []
    for i, it in enumerate(linhas_in, start=1):
        if not isinstance(it, dict):
            erros.append({"linha": i, "erro": "Linha inválida."}); continue
        item_id_raw = it.get("item_id")
        item_id = None
        if item_id_raw not in (None, ""):
            try:
                item_id = int(item_id_raw)
            except (TypeError, ValueError):
                erros.append({"linha": i, "erro": "Item inválido."}); continue
            if item_id not in itens_validos:
                erros.append({"linha": i, "erro": "Item não pertence a esta encomenda."}); continue
        tecido = _texto_celula(it.get("tecido"))
        cor = _texto_celula(it.get("cor"))
        peso = _peso_to_float(it.get("peso"))
        id_rolo = _texto_celula(it.get("id_rolo")) or None
        faltando = [n for n, v in (("tecido", tecido), ("cor", cor)) if not v]
        if faltando:
            erros.append({"linha": i, "erro": f"Campo obrigatório vazio: {', '.join(faltando)}."}); continue
        if peso is None or peso <= 0:
            erros.append({"linha": i, "erro": f"Peso inválido: '{_texto_celula(it.get('peso'))}'."}); continue
        linhas.append({"item_id": item_id, "tecido": tecido, "cor": cor, "peso": round(peso, 3), "id_rolo": id_rolo})
    return linhas, erros

@app.post("/api/encomendas")
@require_login
@require_perm("estoque_mutar")
def criar_encomenda():
    loja = current_loja()
    data = request.get_json(silent=True) or {}
    fornecedor = (data.get("fornecedor") or "").strip() or None
    data_prevista = (data.get("data_prevista") or "").strip() or None
    itens_in = data.get("itens")
    if not isinstance(itens_in, list) or not itens_in:
        return jsonify(error="Informe ao menos um item previsto (tecido, cor, peso_previsto)."), 400

    itens, erros = _validar_itens_encomenda(itens_in)
    if erros:
        return jsonify(error="Há itens inválidos na encomenda.", erros=erros), 400

    conn = get_conn()
    try:
        conn.execute("BEGIN")
        encomenda_id = conn.execute(
            "INSERT INTO encomendas (loja, fornecedor, data_prevista, status) VALUES (?, ?, ?, 'aberta')",
            (loja, fornecedor, data_prevista),
        ).lastrowid
        conn.executemany(
            "INSERT INTO encomenda_itens (encomenda_id, tecido, cor, peso_previsto, rolos_previstos) VALUES (?, ?, ?, ?, ?)",
            [(encomenda_id, it["tecido"], it["cor"], it["peso_previsto"], it["rolos_previstos"]) for it in itens],
        )
        conn.commit()
    except sqlite3.Error as e:
        conn.rollback()
        return jsonify(error=f"Erro de banco de dados: {e}"), 500
    finally:
        conn.close()

    enc, itens_db = _get_encomenda(encomenda_id, loja)
    return jsonify(**enc, itens=itens_db), 201

@app.get("/api/encomendas")
@require_login
@require_perm("estoque_mutar")
def listar_encomendas():
    loja = current_loja()
    status = (request.args.get("status") or "").strip()
    conn = get_conn()
    try:
        query = "SELECT id, fornecedor, data_prevista, status, created_at FROM encomendas WHERE loja = ?"
        params = [loja]
        if status:
            query += " AND status = ?"
            params.append(status)
        query += " ORDER BY id DESC"
        rows = conn.execute(query, params).fetchall()
    finally:
        conn.close()
    return jsonify([dict(r) for r in rows])

@app.get("/api/encomendas/<int:encomenda_id>")
@require_login
@require_perm("estoque_mutar")
def detalhe_encomenda(encomenda_id):
    loja = current_loja()
    enc, itens = _get_encomenda(encomenda_id, loja)
    if not enc:
        return jsonify(error="Encomenda não encontrada."), 404
    return jsonify(**enc, itens=itens)

@app.post("/api/encomendas/<int:encomenda_id>/receber")
@require_login
@require_perm("estoque_mutar")
def receber_encomenda(encomenda_id):
    loja = current_loja()
    enc, itens_previstos = _get_encomenda(encomenda_id, loja)
    if not enc:
        return jsonify(error="Encomenda não encontrada."), 404
    if enc["status"] == "recebida":
        return jsonify(error="Encomenda já recebida."), 409

    data = request.get_json(silent=True) or {}
    fornecedor = (data.get("fornecedor") or "").strip() or enc["fornecedor"]
    data_recebimento = (data.get("data") or "").strip() or datetime.now().strftime("%Y-%m-%d")
    linhas_in = data.get("linhas")
    if not isinstance(linhas_in, list) or not linhas_in:
        return jsonify(error="Informe ao menos uma linha recebida (tecido, cor, peso)."), 400
    if len(linhas_in) > ENTRADA_MAX_LINHAS:
        return jsonify(error=f"Máximo de {ENTRADA_MAX_LINHAS} linhas por recebimento."), 400

    itens_validos = {it["id"] for it in itens_previstos}
    linhas, erros = _validar_linhas_recebimento(linhas_in, itens_validos)
    if erros:
        return jsonify(error="Há linhas inválidas no recebimento.", erros=erros), 400

    ids_rolo = [l["id_rolo"] for l in linhas if l["id_rolo"]]
    repetidos = sorted({r for r in ids_rolo if ids_rolo.count(r) > 1})
    if repetidos:
        return jsonify(error=f"Id de rolo repetido no recebimento: {', '.join(repetidos)}."), 400

    conn = get_conn()
    try:
        conn.execute("BEGIN")
        if ids_rolo:
            marcas = ",".join("?" * len(ids_rolo))
            ja = [r["roll_ext_id"] for r in conn.execute(
                f"SELECT roll_ext_id FROM rolos WHERE loja = ? AND roll_ext_id IN ({marcas})",
                (loja, *ids_rolo),
            ).fetchall()]
            if ja:
                conn.rollback()
                return jsonify(error=f"Rolo(s) já recebido(s) antes: {', '.join(sorted(ja))}."), 409

        tecidos = {
            r["nome_tecido"].casefold(): (r["id"], r["nome_tecido"])
            for r in conn.execute("SELECT id, nome_tecido FROM estoque_tecidos WHERE loja = ?", (loja,)).fetchall()
        }
        desconhecidos = sorted({l["tecido"] for l in linhas if l["tecido"].casefold() not in tecidos})
        if desconhecidos:
            conn.rollback()
            return jsonify(error=f"Tecido(s) não cadastrado(s) no estoque: {', '.join(desconhecidos)}. Cadastre o tecido antes de receber."), 400

        entrada_id = conn.execute(
            "INSERT INTO entradas (loja, fornecedor, data) VALUES (?, ?, ?)",
            (loja, fornecedor, data_recebimento),
        ).lastrowid

        cores_cache = {}
        cores_criadas = []
        for l in linhas:
            tecido_id, tecido_nome = tecidos[l["tecido"].casefold()]
            if tecido_id not in cores_cache:
                cores_cache[tecido_id] = {
                    r["nome_cor"].casefold(): (r["id"], r["nome_cor"])
                    for r in conn.execute("SELECT id, nome_cor FROM estoque_cores WHERE tecido_id = ?", (tecido_id,)).fetchall()
                }
            cores = cores_cache[tecido_id]
            chave = l["cor"].casefold()
            if chave in cores:
                cor_id, cor_nome = cores[chave]
                conn.execute("UPDATE estoque_cores SET peso_kg = ROUND(peso_kg + ?, 3) WHERE id = ?", (l["peso"], cor_id))
            else:
                cor_nome = l["cor"]
                cor_id = conn.execute(
                    "INSERT INTO estoque_cores (tecido_id, nome_cor, peso_kg, qtd_pecas) VALUES (?, ?, ?, 0)",
                    (tecido_id, cor_nome, l["peso"]),
                ).lastrowid
                cores[chave] = (cor_id, cor_nome)
                cores_criadas.append({"tecido": tecido_nome, "cor": cor_nome})
            if l["id_rolo"]:
                conn.execute(
                    "INSERT INTO rolos (loja, entrada_id, tecido, cor, roll_ext_id, peso) VALUES (?, ?, ?, ?, ?, ?)",
                    (loja, entrada_id, tecido_nome, cor_nome, l["id_rolo"], l["peso"]),
                )

        conn.execute("UPDATE encomendas SET status = 'recebida' WHERE id = ?", (encomenda_id,))
        conn.commit()
    except sqlite3.Error as e:
        conn.rollback()
        return jsonify(error=f"Erro de banco de dados: {e}"), 500
    finally:
        conn.close()

    recebido_por_item = {}
    for l in linhas:
        if l["item_id"] is not None:
            recebido_por_item[l["item_id"]] = recebido_por_item.get(l["item_id"], 0.0) + l["peso"]

    variacao = [{
        "item_id": it["id"],
        "tecido": it["tecido"],
        "cor": it["cor"],
        "peso_previsto": it["peso_previsto"],
        "peso_recebido": round(recebido_por_item.get(it["id"], 0.0), 3),
        "variacao": round(recebido_por_item.get(it["id"], 0.0) - it["peso_previsto"], 3),
    } for it in itens_previstos]

    return jsonify(
        id=entrada_id,
        encomenda_id=encomenda_id,
        status="recebida",
        linhas=len(linhas),
        peso_total=round(sum(l["peso"] for l in linhas), 3),
        rolos=len(ids_rolo),
        cores_criadas=cores_criadas,
        variacao=variacao,
    ), 201

# -----------------------------------------------------------------------------
# API: Relatórios (dashboards) — sempre escopados pela loja da sessão
# -----------------------------------------------------------------------------
def _periodo_from_args():
    """Lê ?de=YYYY-MM-DD&ate=YYYY-MM-DD (default: últimos 30 dias).

    Retorna (de, ate, ate_exclusivo): como data_iso é texto "YYYY-MM-DD HH:MM",
    o filtro correto por string é de <= data_iso < ate + 1 dia, cobrindo o
    fim do dia final sem depender de parsing de datas no SQLite.
    """
    def _parse(s):
        try:
            return datetime.strptime((s or "").strip(), "%Y-%m-%d")
        except ValueError:
            return None
    hoje = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
    dt_ate = _parse(request.args.get("ate")) or hoje
    dt_de = _parse(request.args.get("de")) or (dt_ate - timedelta(days=30))
    de = dt_de.strftime("%Y-%m-%d")
    ate = dt_ate.strftime("%Y-%m-%d")
    ate_exclusivo = (dt_ate + timedelta(days=1)).strftime("%Y-%m-%d")
    return de, ate, ate_exclusivo

@app.get("/api/relatorio/loja")
@require_login
@require_perm("relatorio_loja")
def relatorio_loja():
    loja = current_loja()
    # Comissão é DESPESA do dono, não receita da loja: só o admin enxerga
    # o quanto tem a pagar. O gerente vê o desempenho sem esse número.
    ve_comissoes = pode("relatorio_comissoes")
    de, ate, ate_ex = _periodo_from_args()
    conn = get_conn()
    try:
        tot = conn.execute(
            "SELECT COUNT(*) AS n, COALESCE(SUM(total), 0) AS fat "
            "FROM pedidos WHERE loja = ? AND status = 'pedido' AND data_iso >= ? AND data_iso < ?",
            (loja, de, ate_ex),
        ).fetchone()
        num_pedidos = tot["n"]
        faturamento_total = round(tot["fat"], 2)
        ticket_medio = round(faturamento_total / num_pedidos, 2) if num_pedidos else 0.0

        # Só as comissões de quem de fato recebe (o admin nunca recebe).
        comissoes_a_pagar = round(
            conn.execute(
                "SELECT COALESCE(SUM(p.comissao_valor), 0) AS com FROM pedidos p "
                "LEFT JOIN usuarios u ON u.id = p.usuario_id "
                "WHERE p.loja = ? AND p.status = 'pedido' AND p.data_iso >= ? AND p.data_iso < ? "
                "AND COALESCE(u.papel, '') != 'admin'",
                (loja, de, ate_ex),
            ).fetchone()["com"],
            2,
        )

        por_operador = [
            {
                "usuario_id": r["usuario_id"],
                "nome": r["nome"],
                "num_pedidos": r["num_pedidos"],
                "faturamento": round(r["faturamento"], 2),
                **({"comissao": round(r["comissao"], 2)} if ve_comissoes else {}),
            }
            for r in conn.execute(
                "SELECT p.usuario_id, COALESCE(u.nome, '') AS nome, COUNT(*) AS num_pedidos, "
                "COALESCE(SUM(p.total), 0) AS faturamento, COALESCE(SUM(p.comissao_valor), 0) AS comissao "
                "FROM pedidos p LEFT JOIN usuarios u ON u.id = p.usuario_id "
                "WHERE p.loja = ? AND p.status = 'pedido' AND p.data_iso >= ? AND p.data_iso < ? "
                "GROUP BY p.usuario_id ORDER BY faturamento DESC",
                (loja, de, ate_ex),
            ).fetchall()
        ]

        tecidos = [
            {
                "tecido": r["tecido"],
                "faturamento": round(r["faturamento"], 2),
                "peso_total": round(r["peso_total"], 2),
            }
            for r in conn.execute(
                "SELECT p.tecido, COALESCE(SUM(p.total), 0) AS faturamento, "
                "COALESCE((SELECT SUM(i.peso_kg) FROM itens_pedido i "
                "          JOIN pedidos p2 ON p2.id = i.pedido_id "
                "          WHERE p2.loja = ? AND p2.status = 'pedido' AND p2.tecido = p.tecido "
                "            AND p2.data_iso >= ? AND p2.data_iso < ?), 0) AS peso_total "
                "FROM pedidos p "
                "WHERE p.loja = ? AND p.status = 'pedido' AND p.data_iso >= ? AND p.data_iso < ? "
                "GROUP BY p.tecido ORDER BY faturamento DESC",
                (loja, de, ate_ex, loja, de, ate_ex),
            ).fetchall()
        ]

        encalhados = [
            r["nome_tecido"]
            for r in conn.execute(
                "SELECT nome_tecido FROM estoque_tecidos "
                "WHERE loja = ? AND nome_tecido NOT IN ("
                "  SELECT DISTINCT tecido FROM pedidos "
                "  WHERE loja = ? AND status = 'pedido' AND data_iso >= ? AND data_iso < ? AND tecido IS NOT NULL"
                ") ORDER BY nome_tecido",
                (loja, loja, de, ate_ex),
            ).fetchall()
        ]

        if ve_comissoes:
            despesas_total, despesas_por_categoria = _despesas_resumo(conn, loja, de, ate_ex)
    finally:
        conn.close()

    payload = dict(
        periodo={"de": de, "ate": ate},
        faturamento_total=faturamento_total,
        num_pedidos=num_pedidos,
        ticket_medio=ticket_medio,
        por_operador=por_operador,
        tecidos_mais_vendidos=tecidos[:5],
        tecidos_menos_vendidos=sorted(tecidos, key=lambda t: t["faturamento"])[:5],
        tecidos_encalhados=encalhados,
    )
    if ve_comissoes:
        payload["comissoes_a_pagar"] = comissoes_a_pagar
        # Despesas e lucro são números do dono, com a mesma regra das
        # comissões: o gerente não recebe nem a chave no JSON.
        payload["despesas_total"] = despesas_total
        payload["despesas_por_categoria"] = despesas_por_categoria
        payload["lucro"] = round(faturamento_total - despesas_total, 2)
    return jsonify(**payload)

@app.get("/api/relatorio/meu")
@require_login
def relatorio_meu():
    loja = current_loja()
    usuario_id = session.get("usuario_id")
    de, ate, ate_ex = _periodo_from_args()
    conn = get_conn()
    try:
        tot = conn.execute(
            "SELECT COUNT(*) AS n, COALESCE(SUM(total), 0) AS fat, COALESCE(SUM(comissao_valor), 0) AS com "
            "FROM pedidos WHERE loja = ? AND usuario_id = ? AND status = 'pedido' AND data_iso >= ? AND data_iso < ?",
            (loja, usuario_id, de, ate_ex),
        ).fetchone()
        ultimos = [
            {
                "id": r["id"],
                "data_iso": r["data_iso"],
                "cliente_nome": r["cliente_nome"],
                "total": round(r["total"] or 0, 2),
                "comissao_valor": round(r["comissao_valor"] or 0, 2),
            }
            for r in conn.execute(
                "SELECT p.id, p.data_iso, COALESCE(c.nome, '') AS cliente_nome, p.total, p.comissao_valor "
                "FROM pedidos p LEFT JOIN clientes c ON c.id = p.cliente_id "
                "WHERE p.loja = ? AND p.usuario_id = ? AND p.status = 'pedido' AND p.data_iso >= ? AND p.data_iso < ? "
                "ORDER BY p.id DESC LIMIT 10",
                (loja, usuario_id, de, ate_ex),
            ).fetchall()
        ]
    finally:
        conn.close()

    # O admin não recebe comissão: some o número da visão dele em vez de
    # mostrar um zero que parece um erro de cálculo.
    recebe_comissao = papel_atual() != "admin"
    return jsonify(
        periodo={"de": de, "ate": ate},
        num_pedidos=tot["n"],
        faturamento=round(tot["fat"], 2),
        comissao=round(tot["com"], 2) if recebe_comissao else 0.0,
        mostra_comissao=recebe_comissao,
        ultimos_pedidos=ultimos,
    )

# -----------------------------------------------------------------------------
# API: Despesas (somente admin) — importação de planilha xlsx/csv
# -----------------------------------------------------------------------------
PERM_DESPESAS = "despesas_gerir"
DESPESAS_MAX_BYTES = 5 * 1024 * 1024
DESPESAS_MAX_LINHAS = 5000
DESPESAS_COLUNAS = ("data", "descricao", "categoria", "valor")


def _despesas_resumo(conn, loja, de, ate_ex):
    """(total, [{categoria, total}]) das despesas da loja no período."""
    por_categoria = [
        {"categoria": r["categoria"], "total": round(r["total"], 2)}
        for r in conn.execute(
            "SELECT COALESCE(NULLIF(TRIM(categoria), ''), 'Sem categoria') AS categoria, "
            "SUM(valor) AS total FROM despesas "
            "WHERE loja = ? AND data >= ? AND data < ? "
            "GROUP BY 1 ORDER BY total DESC",
            (loja, de, ate_ex),
        ).fetchall()
    ]
    total = round(sum(c["total"] for c in por_categoria), 2)
    return total, por_categoria


def _normaliza_coluna(nome):
    """'Descrição ' -> 'descricao': cabeçalho sem acento, caixa ou espaços."""
    s = unicodedata.normalize("NFKD", str(nome or "")).encode("ascii", "ignore").decode()
    return s.strip().lower()


def _parse_data_despesa(val):
    """Aceita date/datetime (xlsx), 'YYYY-MM-DD', 'DD/MM/YYYY' ou 'DD-MM-YYYY'."""
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, date):
        return val.isoformat()
    s = str(val or "").strip()
    # Datas do Excel exportadas em CSV às vezes vêm com a hora junto.
    s = s.split(" ")[0].split("T")[0]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def _parse_valor_despesa(val):
    """Número da planilha em pt-BR ('1.234,56', 'R$ 45,90') ou ponto decimal ('45.90').

    Diferente de _ptbr_to_float, não descarta o ponto às cegas: CSVs gerados
    por sistemas costumam usar '45.90', que viraria 4590.
    """
    if isinstance(val, bool) or val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = re.sub(r"[^\d,.\-]", "", str(val).strip())
    if not s:
        return None
    if "," in s and "." in s:
        # O separador que aparece por último é o decimal.
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        # '1.500' em planilha brasileira é mil e quinhentos.
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return None


def _linhas_csv(conteudo):
    try:
        texto = conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Excel em português salva CSV em Windows-1252.
        texto = conteudo.decode("cp1252", errors="replace")
    try:
        dialeto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t")
    except csv.Error:
        dialeto = csv.excel
    return list(csv.reader(io.StringIO(texto), dialeto))


def _linhas_xlsx(conteudo):
    wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    try:
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()


@app.post("/api/despesas/importar")
@require_login
@require_perm(PERM_DESPESAS)
def despesas_importar():
    loja = current_loja()
    arquivo = request.files.get("arquivo")
    if not arquivo or not arquivo.filename:
        return jsonify(error="Envie a planilha no campo 'arquivo'."), 400
    nome = arquivo.filename.lower()
    conteudo = arquivo.read(DESPESAS_MAX_BYTES + 1)
    if len(conteudo) > DESPESAS_MAX_BYTES:
        return jsonify(error="Planilha muito grande (máx. 5 MB)."), 400

    try:
        if nome.endswith(".csv"):
            linhas = _linhas_csv(conteudo)
        elif nome.endswith(".xlsx"):
            linhas = _linhas_xlsx(conteudo)
        else:
            return jsonify(error="Formato não suportado: use .xlsx ou .csv."), 400
    except Exception:
        return jsonify(error="Não foi possível ler a planilha. Verifique o arquivo."), 400

    if not linhas:
        return jsonify(error="Planilha vazia."), 400
    cabecalho = [_normaliza_coluna(c) for c in linhas[0]]
    faltando = [c for c in ("data", "valor") if c not in cabecalho]
    if faltando:
        return jsonify(
            error=f"Coluna(s) obrigatória(s) ausente(s): {', '.join(faltando)}. "
                  f"Esperado: {', '.join(DESPESAS_COLUNAS)}."
        ), 400
    idx = {c: cabecalho.index(c) for c in DESPESAS_COLUNAS if c in cabecalho}
    if len(linhas) - 1 > DESPESAS_MAX_LINHAS:
        return jsonify(error=f"Máximo de {DESPESAS_MAX_LINHAS} linhas por importação."), 400

    def _cel(row, col):
        i = idx.get(col)
        return row[i] if i is not None and i < len(row) else None

    validas, erros = [], []
    # Linha 1 é o cabeçalho: numeramos como o usuário vê na planilha.
    for n, row in enumerate(linhas[1:], start=2):
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        data_raw, valor_raw = _cel(row, "data"), _cel(row, "valor")
        data_iso = _parse_data_despesa(data_raw)
        valor = _parse_valor_despesa(valor_raw)
        problemas = []
        if not data_iso:
            problemas.append(f"data inválida ({data_raw!s})" if data_raw not in (None, "") else "data vazia")
        if valor is None:
            problemas.append(f"valor inválido ({valor_raw!s})" if valor_raw not in (None, "") else "valor vazio")
        elif valor <= 0:
            problemas.append("valor deve ser maior que zero")
        if problemas:
            erros.append({"linha": n, "erro": "; ".join(problemas)})
            continue
        descricao = str(_cel(row, "descricao") or "").strip()
        categoria = str(_cel(row, "categoria") or "").strip()
        validas.append((loja, data_iso, descricao, categoria, round(valor, 2)))

    conn = get_conn()
    try:
        conn.executemany(
            "INSERT INTO despesas (loja, data, descricao, categoria, valor) VALUES (?, ?, ?, ?, ?)",
            validas,
        )
        conn.commit()
    except sqlite3.Error as e:
        conn.rollback()
        return jsonify(error=f"Erro de banco de dados: {e}"), 500
    finally:
        conn.close()

    return jsonify(
        importadas=len(validas),
        erros=erros,
        total=round(sum(v[4] for v in validas), 2),
    ), 200


@app.get("/api/despesas")
@require_login
@require_perm(PERM_DESPESAS)
def despesas_listar():
    loja = current_loja()
    de, ate, ate_ex = _periodo_from_args()
    conn = get_conn()
    try:
        despesas = [dict(r) for r in conn.execute(
            "SELECT id, data, descricao, categoria, valor FROM despesas "
            "WHERE loja = ? AND data >= ? AND data < ? ORDER BY data DESC, id DESC",
            (loja, de, ate_ex),
        ).fetchall()]
        total, por_categoria = _despesas_resumo(conn, loja, de, ate_ex)
    finally:
        conn.close()
    return jsonify(periodo={"de": de, "ate": ate}, despesas=despesas,
                   total=total, por_categoria=por_categoria)


@app.delete("/api/despesas/<int:despesa_id>")
@require_login
@require_perm(PERM_DESPESAS)
def despesa_delete(despesa_id):
    loja = current_loja()
    conn = get_conn()
    try:
        cur = conn.execute("DELETE FROM despesas WHERE id = ? AND loja = ?", (despesa_id, loja))
        if cur.rowcount == 0:
            return jsonify(error="Despesa não encontrada."), 404
        conn.commit()
    finally:
        conn.close()
    return jsonify(ok=True)

# -----------------------------------------------------------------------------
# API: PIX e Exportações
# -----------------------------------------------------------------------------
@app.post("/pix")
@require_login
@require_perm("config_editar")
def pix_salvar():
    loja = current_loja()
    data = request.get_json(silent=True) or {}
    pix_chave = (data.get("pix_chave") or "").strip()
    if callable(salvar_pix_by_codigo):
        salvar_pix_by_codigo(codigo=loja, pix_chave=pix_chave)
        return jsonify(ok=True, message="Chave PIX salva com sucesso.")
    return jsonify(error="Função de salvar PIX indisponível."), 500

@app.get("/exportar/<int:pedido_id>")
@require_login
def exportar_pedido(pedido_id: int):
    loja_codigo = current_loja()
    tipo = (request.args.get("type") or "png").lower()
    pedido, itens = _carregar_pedido_itens(pedido_id, loja_codigo)
    if not pedido:
        return jsonify(error="Pedido não encontrado ou não pertence à sua loja."), 404
    loja_info = get_loja_by_codigo(loja_codigo) or {}
    if tipo == "png":
        if not callable(render_pedido_png):
            return jsonify(error="Serviço de renderização de PNG indisponível."), 500
        buf = render_pedido_png(loja=loja_info, pedido=pedido, itens=itens)
        return send_file(buf, mimetype="image/png", as_attachment=True, download_name=f"romaneio_{pedido_id}.png")
    elif tipo == "pdf":
        if not callable(render_pedido_pdf):
            return jsonify(error="Serviço de renderização de PDF indisponível."), 500
        with app.app_context():
            buf = render_pedido_pdf(loja=loja_info, pedido=pedido, itens=itens)
        return send_file(buf, mimetype="application/pdf", as_attachment=True, download_name=f"romaneio_{pedido_id}.pdf")
    return jsonify(error="Tipo de exportação inválido."), 400

# -----------------------------------------------------------------------------
# Rotas de Admin
# -----------------------------------------------------------------------------
@app.get("/admin/lojas")
@require_admin
def admin_get_todas_lojas():
    conn = get_conn()
    lojas = [dict(r) for r in conn.execute("SELECT codigo, nome, status, created_at FROM lojas ORDER BY id DESC").fetchall()]
    conn.close(); return jsonify(lojas)

@app.post("/admin/aprovar-loja/<codigo_loja>")
@require_admin
def admin_aprovar_loja(codigo_loja):
    conn = get_conn()
    conn.execute("UPDATE lojas SET status = 'aprovado' WHERE codigo = ?", (codigo_loja,))
    conn.commit(); conn.close()
    return jsonify(ok=True, message=f"Loja {codigo_loja} aprovada.")

@app.post("/admin/revogar-loja/<codigo_loja>")
@require_admin
def admin_revogar_loja(codigo_loja):
    conn = get_conn()
    conn.execute("UPDATE lojas SET status = 'revogado' WHERE codigo = ?", (codigo_loja,))
    conn.commit(); conn.close()
    return jsonify(ok=True, message=f"Acesso da loja {codigo_loja} foi revogado.")

@app.post("/admin/reativar-loja/<codigo_loja>")
@require_admin
def admin_reativar_loja(codigo_loja):
    conn = get_conn()
    conn.execute("UPDATE lojas SET status = 'aprovado' WHERE codigo = ?", (codigo_loja,))
    conn.commit(); conn.close()
    return jsonify(ok=True, message=f"Acesso da loja {codigo_loja} foi reativado.")

@app.post("/admin/deletar-loja/<codigo_loja>")
@require_admin
def admin_deletar_loja(codigo_loja):
    conn = get_conn()
    try:
        cur = conn.execute("DELETE FROM lojas WHERE codigo = ?", (codigo_loja,))
        if cur.rowcount == 0:
            conn.close(); return jsonify(error="Loja não encontrada."), 404
        conn.commit()
    except sqlite3.Error as e:
        conn.rollback(); return jsonify(error=f"Erro no banco de dados: {e}"), 500
    finally: conn.close()
    return jsonify(ok=True, message=f"Loja {codigo_loja} deletada com sucesso.")

# -----------------------------------------------------------------------------
# Main
# -----------------------------------------------------------------------------
if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5001))
    app.run(debug=False, host="0.0.0.0", port=port)

